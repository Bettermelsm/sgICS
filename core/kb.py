"""知识库：文档解析 → 切分 → embedding → 检索。

设计取舍（v0.7）：
- embedding 走智谱 embedding 接口（与 GLM 同一 Key）；
- 向量直接存 DB（JSON），Python 暴力算余弦相似度。
  demo 量级（几百上千块）完全够用；上万块以后再迁向量数据库。
"""
import hashlib
import json
import math
import os

CHUNK_SIZE = 500
CHUNK_OVERLAP = 50
EMBED_BATCH = 32


# ---------- 文档解析 ----------
def extract_text(path: str, doc_type: str) -> str:
    doc_type = (doc_type or '').lower()
    if doc_type in ('txt', 'md', ''):
        with open(path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
    if doc_type == 'pdf':
        from pypdf import PdfReader
        return '\n'.join((p.extract_text() or '') for p in PdfReader(path).pages)
    if doc_type == 'docx':
        from docx import Document
        return '\n'.join(p.text for p in Document(path).paragraphs)
    if doc_type == 'xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                line = ' | '.join(str(c).strip() for c in row if c not in (None, ''))
                if line:
                    lines.append(line)
        return '\n'.join(lines)
    raise ValueError(f'不支持的文档类型：{doc_type}（支持 txt/md/pdf/docx/xlsx）')


def guess_doc_type(filename: str) -> str:
    name = (filename or '').lower()
    for ext in ('txt', 'md', 'pdf', 'docx', 'xlsx'):
        if name.endswith('.' + ext):
            return ext
    return ''


# ---------- 切分 ----------
def chunk_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list:
    paras = [p.strip() for p in text.splitlines() if p.strip()]
    # 超长段落先硬切（保证每块 ≤ size）
    units = []
    for p in paras:
        while len(p) > size:
            units.append(p[:size])
            p = p[size - overlap:] if overlap else p[size:]
        if p:
            units.append(p)
    chunks, buf = [], ''
    for u in units:
        cand = f'{buf}\n{u}'.strip() if buf else u
        if len(cand) > size and buf:
            chunks.append(buf)
            buf = buf[-overlap:] if overlap and len(buf) > overlap else ''
            cand = f'{buf}\n{u}'.strip() if buf else u
            if len(cand) > size:
                buf, cand = '', u  # overlap 后仍超：放弃 overlap
        buf = cand
    if buf:
        chunks.append(buf)
    return [c for c in chunks if c]


# ---------- embedding ----------
def embed_texts(texts: list) -> list:
    """智谱 embedding 接口（OpenAI 兼容）。需要 ZHIPU_API_KEY。"""
    if not texts:
        return []
    from openai import OpenAI
    api_key = os.environ.get('ZHIPU_API_KEY', '')
    if not api_key:
        raise RuntimeError('缺少 ZHIPU_API_KEY：知识库入库需要调用 embedding 接口')
    base_url = os.environ.get('ZHIPU_BASE_URL', 'https://open.bigmodel.cn/api/paas/v4/')
    model = os.environ.get('ZHIPU_EMBED_MODEL', 'embedding-2')
    client = OpenAI(api_key=api_key, base_url=base_url, timeout=60)
    vectors = []
    for i in range(0, len(texts), EMBED_BATCH):
        resp = client.embeddings.create(model=model, input=texts[i:i + EMBED_BATCH])
        vectors.extend([d.embedding for d in resp.data])
    return vectors


def fake_embed(texts: list) -> list:
    """确定性 fake embedding（测试用）：按字符哈希到 64 维并归一化。"""
    vecs = []
    for t in texts:
        v = [0.0] * 64
        for ch in t:
            h = int(hashlib.md5(ch.encode()).hexdigest(), 16)
            v[h % 64] += 1.0
        n = math.sqrt(sum(x * x for x in v)) or 1.0
        vecs.append([x / n for x in v])
    return vecs


# ---------- 检索 ----------
def cosine(a: list, b: list) -> float:
    denom = math.sqrt(sum(x * x for x in a)) * math.sqrt(sum(x * x for x in b))
    if not denom:
        return 0.0
    return sum(x * y for x, y in zip(a, b)) / denom


class Retriever:
    """某店铺的知识库检索器。"""

    def __init__(self, shop, embed_fn=None, min_score: float = 0.25):
        self.shop = shop
        self.embed_fn = embed_fn or embed_texts
        self.min_score = min_score

    def _vectors(self):
        from .models import KnowledgeChunk
        out = []
        qs = (KnowledgeChunk.objects
              .filter(doc__shop=self.shop, doc__status='ready')
              .exclude(embedding=''))
        for c in qs:
            try:
                out.append((c, json.loads(c.embedding)))
            except Exception:
                continue
        return out

    def search(self, query: str, top_k: int = 3) -> list:
        """返回 [(chunk, score)]，按相似度降序。"""
        qv = self.embed_fn([query])[0]
        scored = [(c, cosine(qv, v)) for c, v in self._vectors()]
        scored = [(c, s) for c, s in scored if s >= self.min_score]
        scored.sort(key=lambda x: x[1], reverse=True)
        return scored[:top_k]


# ---------- 入库 ----------
def ingest_doc(doc, embed_fn=None) -> int:
    """解析→切分→向量化→入库。成功返回分块数；失败置 failed 并抛异常。"""
    from .models import KnowledgeChunk
    doc.status = 'indexing'
    doc.save(update_fields=['status'])
    try:
        text = extract_text(doc.file.path, doc.doc_type)
        chunks = chunk_text(text)
        if not chunks:
            raise ValueError('文档解析为空，没有可入库的内容')
        vectors = (embed_fn or embed_texts)(chunks)
        KnowledgeChunk.objects.filter(doc=doc).delete()
        objs = [KnowledgeChunk(doc=doc, ordering=i, content=ch,
                               embedding=json.dumps(v))
                for i, (ch, v) in enumerate(zip(chunks, vectors))]
        KnowledgeChunk.objects.bulk_create(objs)
        doc.chunk_count = len(objs)
        doc.status = 'ready'
        doc.error = ''
        doc.save(update_fields=['chunk_count', 'status', 'error'])
        return len(objs)
    except Exception as e:
        doc.status = 'failed'
        doc.error = str(e)[:500]
        doc.save(update_fields=['status', 'error'])
        raise
